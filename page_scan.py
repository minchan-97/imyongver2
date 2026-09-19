"""
page_scan.py — 파일을 '페이지 통째로' 스캔해서 읽는다.

기존(file_ingest): 텍스트 추출 → 문장 단위로 잘게 분할 → 사람이 한 줄씩 검토
  → 느리고, 표·2단 PDF는 문장이 섞이고, 스캔본은 못 읽음.

지금: PDF 페이지를 이미지로 렌더 → 비전 모델이 페이지 전체를 읽음
      (표·레이아웃 유지한 본문 + 성취기준 코드 + 영역 + 개념 태그를 한 번에)
  - 1페이지 = 1기록. 사람은 '틀린 페이지만' 고친다.
  - 여러 페이지 동시 처리(workers).
  - 결과는 (파일해시, 페이지)로 캐시 → 같은 파일 재업로드 시 재스캔 없음.
    캐시는 data/scan_cache.pkl → 서버에도 자동 백업.
  - 키가 없으면: PDF 텍스트층으로 대체(스캔본은 건너뜀 표시).
  - docx/txt는 페이지가 없으니 일정 길이로 묶어 '페이지'로 취급(LLM 불필요).
"""
from __future__ import annotations
import io, os, re, json, time, base64, pickle, hashlib
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    from cloud import push as _cloud_push
except Exception:
    def _cloud_push(path, **kw): return False

CODE_RE = re.compile(r"\[\d{1,2}[가-힣]{1,3}\d{2}[-–]\d{2}\]")
DEFAULT_MODEL = "gpt-4o-mini"

SYSTEM = """너는 한국 초등 임용시험 자료(교육과정·교사용 지도서·기출 시험지) 스캔 페이지를 읽는 판독기다.
페이지 이미지를 보고 아래 JSON 하나만 출력한다. 설명·코드블록 금지.
{
 "text": "페이지 본문 전체. 읽는 순서대로. 표는 마크다운 표로, 2단 편집은 왼쪽 단 → 오른쪽 단 순서. 쪽번호·머리글·바닥글은 제외.",
 "codes": ["페이지에 실제로 적힌 성취기준 코드들, 예: [4국02-01]"],
 "area": "영역 한 단어(듣기·말하기/읽기/쓰기/문법/문학/매체 등). 판단 불가면 빈 문자열",
 "concepts": ["주어진 개념 목록 중 이 페이지가 직접 다루는 것만. 목록에 없는 말은 만들지 말 것"],
 "skip": false
}
표지·목차·빈 페이지·판권면이면 skip을 true로.
본문은 요약하지 말고 적힌 그대로 옮긴다. 안 보이는 글자는 추측하지 말고 [판독불가]로 둔다."""

HANDWRITING = """
이 페이지에는 한글 손글씨(필기·노트·여백 메모)가 있을 수 있다.
- 손글씨도 인쇄 글자와 똑같이 한 글자씩 옮긴다. 맞춤법을 고치거나 문장을 다듬지 말 것.
- 흘려 쓴 글자는 앞뒤 문맥(교육과정·국어교육 용어)으로 판독하되, 확신이 없으면 그 낱말 뒤에 (?)를 붙인다.
- 줄 친 곳·동그라미·별표는 해당 부분 앞에 [강조]를 붙인다.
- 화살표로 연결된 내용은 'A → B'로 적는다.
- 여백 메모는 해당 위치에 (메모: …)로 끼워 넣는다.
- 인쇄된 본문 위에 필기가 겹쳐 있으면 인쇄 본문을 먼저, 필기를 (필기: …)로 뒤에 적는다."""

HANDWRITING_MODEL = "gpt-4o"   # 손글씨는 mini보다 확실히 나음


def file_sha(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


# ── 캐시 ─────────────────────────────────────────────────────
class ScanCache:
    def __init__(self, path):
        self.path = path
        self.d = {}
        if os.path.exists(path):
            try:
                with open(path, "rb") as f:
                    self.d = pickle.load(f)
            except Exception:
                self.d = {}

    def get(self, sha, page, mode):
        return self.d.get(f"{sha}|{page}|{mode}")

    def put(self, sha, page, mode, result):
        self.d[f"{sha}|{page}|{mode}"] = result

    def save(self):
        with open(self.path, "wb") as f:
            pickle.dump(self.d, f)
        _cloud_push(self.path)


# ── 페이지 준비: 여러 파일 → '쪽' 목록 ─────────────────────
IMAGE_EXT = (".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif")

def _is_image(name): return name.lower().endswith(IMAGE_EXT)
def _is_pdf(name): return name.lower().endswith(".pdf")


def _pdf_doc(raw):
    import pypdfium2 as pdfium
    return pdfium.PdfDocument(raw)


def sniff(name: str, raw: bytes) -> str:
    """
    확장자 대신 '실제 내용'으로 종류 판별.
    폰에서 올린 파일은 이름과 속이 다른 경우가 많다
    (HWP를 .pdf로 저장, 다운로드 실패로 받은 웹페이지, PNG인데 .pdf 등).
    """
    head = raw[:2048]
    if b"%PDF" in head[:1024]:
        return "pdf"
    if head[:8] == b"\x89PNG\r\n\x1a\n" or head[:3] == b"\xff\xd8\xff" \
            or head[:4] == b"RIFF" or head[4:12] in (b"ftypheic", b"ftypheix",
                                                    b"ftypmif1", b"ftyphevc"):
        return "image"
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or b"HWP Document File" in head:
        return "hwp"
    if head[:4] == b"PK\x03\x04":
        return "hwpx" if name.lower().endswith(".hwpx") else "docx"
    low = head.lower()
    if b"<html" in low or b"<!doctype html" in low:
        return "html"
    if name.lower().endswith(IMAGE_EXT):
        return "image"
    return "text"


BAD_KIND_MSG = {
    "hwp": "한글(HWP) 파일이에요. 한글에서 'PDF로 저장' 후 올려주세요.",
    "hwpx": "한글(HWPX) 파일이에요. 한글에서 'PDF로 저장' 후 올려주세요.",
    "html": "PDF가 아니라 웹페이지(다운로드 실패 화면)예요. 원본 PDF를 다시 받아주세요.",
}


def _pdf_open_error(e) -> str:
    m = str(e)
    if "password" in m.lower():
        return "암호가 걸린 PDF예요. 암호를 풀어 다시 저장한 뒤 올려주세요."
    return f"PDF를 열 수 없어요(손상된 파일일 수 있음): {m}"


def build_pages(files) -> list[dict]:
    """
    files: [(파일명, bytes), ...] (업로드 순서 유지)
    반환: 쪽 목록 [{name, raw, sha, index, kind, text?, error?}]
      pdf → 쪽마다 1개 / 사진 → 파일마다 1개 / docx·txt → 약 1800자씩 1개
      열 수 없는 파일 → kind="broken" 1개 (앱이 죽지 않고 그 파일만 오류 표시)
    """
    pages = []
    for name, raw in files:
        sha = file_sha(raw)
        base = {"name": name, "raw": raw, "sha": sha}
        kind = sniff(name, raw)
        try:
            if kind in BAD_KIND_MSG:
                raise ValueError(BAD_KIND_MSG[kind])
            if kind == "pdf":
                start = raw.find(b"%PDF")          # 앞에 붙은 쓰레기 바이트 제거
                if start > 0:
                    raw = raw[start:]
                    base["raw"] = raw
                try:
                    n = len(_pdf_doc(raw))
                except Exception as e:
                    # pdfium이 못 열면 pdfplumber로 텍스트라도 건짐
                    texts = _pdfplumber_texts(raw)
                    if texts is None:
                        raise ValueError(_pdf_open_error(e))
                    for i, t in enumerate(texts):
                        pages.append(dict(base, index=i, kind="text", text=t))
                    continue
                if n == 0:
                    raise ValueError("쪽이 0개인 PDF예요.")
                for i in range(n):
                    pages.append(dict(base, index=i, kind="pdf"))
            elif kind == "image":
                pages.append(dict(base, index=0, kind="image"))
            else:
                tp = _text_pages(name if kind != "docx" else "x.docx", raw)
                if not tp:
                    raise ValueError("읽을 글자가 없는 파일이에요.")
                for i, t in enumerate(tp):
                    pages.append(dict(base, index=i, kind="text", text=t))
        except Exception as e:
            pages.append(dict(base, index=0, kind="broken", error=str(e)))
    return pages


def _pdfplumber_texts(raw):
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            texts = [(p.extract_text() or "") for p in pdf.pages]
        return texts if any(t.strip() for t in texts) else None
    except Exception:
        return None


def _to_jpeg(img, max_side=2000) -> bytes:
    from PIL import ImageOps
    img = ImageOps.exif_transpose(img).convert("RGB")   # 폰 사진 회전 보정
    w, h = img.size
    if max(w, h) > max_side:
        k = max_side / max(w, h)
        img = img.resize((int(w * k), int(h * k)))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def render_page(raw: bytes, index: int, dpi: int = 110) -> bytes:
    """PDF 한 쪽 → JPEG."""
    pdf = _pdf_doc(raw)
    return _to_jpeg(pdf[index].render(scale=dpi / 72).to_pil())


def render_source(page: dict, dpi: int = 110):
    """쪽 하나 → 비전 입력/미리보기용 JPEG (텍스트 쪽은 None)."""
    if page["kind"] == "pdf":
        return render_page(page["raw"], page["index"], dpi)
    if page["kind"] == "broken":
        return None
    if page["kind"] == "image":
        from PIL import Image
        try:
            import pillow_heif                     # 아이폰 HEIC 지원(설치돼 있으면)
            pillow_heif.register_heif_opener()
        except Exception:
            pass
        try:
            return _to_jpeg(Image.open(io.BytesIO(page["raw"])))
        except Exception as e:
            raise ValueError(f"사진을 열 수 없어요(HEIC면 pillow-heif 필요): {e}")
    return None


def _pdf_text_layer(raw: bytes, index: int) -> str:
    pdf = _pdf_doc(raw)
    return pdf[index].get_textpage().get_text_range() or ""


def _text_pages(filename, raw, size=1800) -> list[str]:
    name = filename.lower()
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(raw))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
    else:
        paras = [p for p in raw.decode("utf-8", errors="ignore").split("\n") if p.strip()]
    pages, cur = [], ""
    for p in paras:
        if cur and len(cur) + len(p) > size:
            pages.append(cur)
            cur = ""
        cur += p + "\n"
    if cur.strip():
        pages.append(cur)
    return pages


# ── 한 쪽 스캔 ────────────────────────────────────────────────
def _codes(text):
    return sorted(set(CODE_RE.findall(text)))


def _scan_vision(jpeg: bytes, api_key: str, model: str, concept_names,
                 handwriting=False) -> dict:
    from openai import OpenAI
    client = OpenAI(api_key=api_key)
    system = SYSTEM + (HANDWRITING if handwriting else "")
    hint = "개념 목록: " + (", ".join(concept_names) if concept_names else "(없음)")
    b64 = base64.b64encode(jpeg).decode()
    last = None
    for attempt in range(3):
        try:
            r = client.chat.completions.create(
                model=model, temperature=0,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": [
                        {"type": "text", "text": hint},
                        {"type": "image_url", "image_url": {
                            "url": f"data:image/jpeg;base64,{b64}", "detail": "high"}},
                    ]},
                ])
            d = json.loads(r.choices[0].message.content)
            text = str(d.get("text", "")).strip()
            allowed = set(concept_names or [])
            return {
                "text": text,
                "codes": sorted(set(d.get("codes") or []) | set(_codes(text))),
                "area": str(d.get("area") or "").strip(),
                "concepts": [c for c in (d.get("concepts") or []) if c in allowed],
                "skip": bool(d.get("skip")) or not text,
                "method": "vision-hw" if handwriting else "vision",
            }
        except Exception as e:
            last = e
            time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"스캔 실패: {last}")


def _plain(text, method):
    text = (text or "").strip()
    return {"text": text, "codes": _codes(text), "area": "", "concepts": [],
            "skip": len(text) < 20, "method": method}


# ── 전체 스캔 ─────────────────────────────────────────────────
def scan_pages(pages, api_key=None, concept_names=None, cache: ScanCache = None,
               workers=6, model=None, handwriting=False, progress=None) -> list[dict]:
    """
    pages: build_pages() 결과
    반환: 쪽별 dict {page(0부터 전체 순번), name, index, text, codes, area,
                     concepts, skip, method, cached, error?, note?}
    handwriting=True → 손글씨 지침 추가 + 고해상도 + 손글씨용 모델(기본 gpt-4o)
    """
    if handwriting:
        model = model or HANDWRITING_MODEL
    model = model or DEFAULT_MODEL
    dpi = 160 if handwriting else 110
    mode = f"vision:{model}:{'hw' if handwriting else 'print'}" if api_key else "textlayer"
    n = len(pages)
    results = [None] * n
    todo = []
    for k, pg in enumerate(pages):
        meta = {"page": k, "name": pg["name"], "index": pg["index"]}
        if pg["kind"] == "broken":
            results[k] = dict(_plain("", "broken"), cached=False, error=pg["error"], **meta)
            continue
        if pg["kind"] == "text":
            results[k] = dict(_plain(pg["text"], "text"), cached=False, **meta)
            continue
        hit = cache.get(pg["sha"], pg["index"], mode) if cache else None
        if hit:
            results[k] = dict(hit, cached=True, **meta)
        else:
            todo.append(k)

    done = n - len(todo)
    if progress:
        progress(done, n)

    def work(k):
        pg = pages[k]
        if api_key:
            return _scan_vision(render_source(pg, dpi), api_key, model, concept_names,
                                handwriting)
        if pg["kind"] == "pdf":
            r = _plain(_pdf_text_layer(pg["raw"], pg["index"]), "textlayer")
            if not r["text"]:
                r["note"] = "텍스트층 없음(스캔본) — OpenAI 키가 있으면 읽을 수 있어요"
            return r
        r = _plain("", "none")
        r["note"] = "사진은 OpenAI 키가 있어야 읽을 수 있어요"
        return r

    if todo:
        with ThreadPoolExecutor(max_workers=workers if api_key else 1) as ex:
            futs = {ex.submit(work, k): k for k in todo}
            for fut in as_completed(futs):
                k = futs[fut]
                pg = pages[k]
                meta = {"page": k, "name": pg["name"], "index": pg["index"]}
                try:
                    r = fut.result()
                    if cache and r.get("method") != "none":
                        cache.put(pg["sha"], pg["index"], mode, r)
                    results[k] = dict(r, cached=False, **meta)
                except Exception as e:
                    results[k] = dict(_plain("", "error"), cached=False, error=str(e), **meta)
                done += 1
                if progress:
                    progress(done, n)
        if cache:
            cache.save()
    return results


def scan_file(filename, raw, api_key=None, concept_names=None, cache=None,
              workers=6, model=None, handwriting=False, progress=None):
    """파일 1개용 (호환)."""
    return scan_pages(build_pages([(filename, raw)]), api_key, concept_names, cache,
                      workers, model, handwriting, progress)
